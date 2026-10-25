"""Pair sampling, workbook and statistics of the deduplication audit (no data files read)."""
from __future__ import annotations

import hashlib
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
from dedup_audit_agreement import (analyse, clopper_pearson, disagreements,  # noqa: E402
                                   human_vs_llm, stratum_summary)
from make_dedup_audit_workbook import (DECISIONS, FIELDS, RawRecords, assemble, clip,  # noqa: E402
                                       doi_cluster_id, format_authors, sample_judged_pairs,
                                       sample_layer1_pairs, write_workbook)
from second_reviewer_agreement import validate_decisions  # noqa: E402

SAME, DIFF, UND = DECISIONS


def cid(key: str) -> str:
    return hashlib.sha1(key.encode()).hexdigest()[:16]


def synthetic_full() -> pd.DataFrame:
    """10 singletons, 6 DOI clusters of size 2-3 and 4 title clusters of size 2."""
    rows = []
    for i in range(10):
        rows.append(dict(cluster_id=f"single{i}", is_primary=True, source="wos", source_id=f"s{i}", doi=None))
    for c in range(6):
        doi = f"10.1/{c}"
        for m in range(2 + c % 2):
            rows.append(dict(cluster_id=cid(f"doi::{doi}"), is_primary=m == 0, source="scopus",
                             source_id=f"d{c}_{m}", doi=doi))
    for c in range(4):
        for m in range(2):
            rows.append(dict(cluster_id=cid(f"title::t{c}::2020"), is_primary=m == 1, source="scielo",
                             source_id=f"t{c}_{m}", doi=None))
    return pd.DataFrame(rows)


def test_format_authors_and_clip():
    assert format_authors([f"A{i}" for i in range(8)]) == "A0; A1; A2; A3; A4; A5 et al."
    assert format_authors(np.array(["X", "Y"])) == "X; Y"
    assert format_authors(None) == "" and format_authors([]) == ""
    assert clip("a" * 500).endswith("…") and len(clip("a" * 500)) == 401
    assert clip("corto  con\nespacios") == "corto con espacios"
    assert clip(None) == "" and clip(float("nan")) == ""
    assert clip("evaluÃ³") == "evaluó"


def test_doi_cluster_id_matches_deterministic_rule():
    assert doi_cluster_id("10.1/0") == cid("doi::10.1/0")
    assert doi_cluster_id(None) is None and doi_cluster_id("") is None


def test_layer1_sampling_uses_distinct_multi_record_clusters():
    full = synthetic_full()
    pairs = sample_layer1_pairs(full, 10, np.random.default_rng(1))
    assert len(pairs) == 10 and pairs.cluster_id.is_unique
    assert not pairs.cluster_id.str.startswith("single").any()
    assert full.loc[pairs.a_pos, "is_primary"].all()
    assert not full.loc[pairs.b_pos, "is_primary"].any()
    assert (full.loc[pairs.a_pos, "cluster_id"].to_numpy() == full.loc[pairs.b_pos, "cluster_id"].to_numpy()).all()
    assert pairs.match_basis.value_counts().to_dict() == {"doi": 6, "title_year": 4}
    again = sample_layer1_pairs(full, 10, np.random.default_rng(1))
    pd.testing.assert_frame_equal(pairs, again)
    with pytest.raises(ValueError):
        sample_layer1_pairs(full, 11, np.random.default_rng(1))


def test_judged_sampling_quotas_take_all_when_short():
    judged = pd.DataFrame({"judge_verdict": ["DUPLICATE"] * 50 + ["NOT_DUPLICATE"] * 40 + ["UNCLEAR"] * 7,
                           "score": np.linspace(0.9, 1, 97)})
    out = sample_judged_pairs(judged, {"DUPLICATE": 30, "NOT_DUPLICATE": 30, "UNCLEAR": 10},
                              np.random.default_rng(2))
    assert out.judge_verdict.value_counts().to_dict() == {"DUPLICATE": 30, "NOT_DUPLICATE": 30, "UNCLEAR": 7}
    assert out.index.is_unique


def side_frame(n: int, stratum: str, tag: str = "x") -> pd.DataFrame:
    d = {f"{s}_{f}": [f"{tag}{i}-{s}-{f}" for i in range(n)] for s in "ab" for f in FIELDS + ["ref"]}
    return pd.DataFrame({**d, "stratum": stratum})


def test_assemble_shuffles_and_swaps_consistently():
    pairs = assemble(side_frame(40, "L1", "x"), side_frame(30, "L3", "y"), np.random.default_rng(3))
    assert pairs.n.tolist() == list(range(1, 71))
    assert pairs.stratum.value_counts().to_dict() == {"L1": 40, "L3": 30}
    assert pairs.stratum.iloc[:40].nunique() == 2              # strata are mixed
    for r in pairs.itertuples():
        side = "b" if r.swapped else "a"
        for f in FIELDS + ["ref"]:                              # all fields of a side move together
            assert getattr(r, f"a_{f}").endswith(f"-{side}-{f}")
    assert 0 < pairs.swapped.sum() < 70


def test_workbook_is_blind_and_has_dropdown(tmp_path):
    pairs = assemble(side_frame(3, "L1_merge", "x"), side_frame(2, "L3_DUPLICATE", "y"),
                     np.random.default_rng(4))
    for s in "ab":
        pairs[f"{s}_year"] = 2020
    out = tmp_path / "audit.xlsx"
    write_workbook(pairs, out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Instrucciones", "Pares"]
    ws = wb["Pares"]
    dv = ws.data_validations.dataValidation
    assert len(dv) == 1 and str(dv[0].sqref) == "B2:B6" and "Mismo trabajo" in dv[0].formula1
    header = [c.value for c in ws[1]]
    assert header[:2] == ["N°", "Decisión"] and header[-1] == "Comentario"
    values = " ".join(str(c.value) for row in ws.iter_rows(min_row=2) for c in row if c.value is not None)
    assert "L1_merge" not in values and "L3_DUPLICATE" not in values     # stratum never written
    counter = [c.value for c in wb["Instrucciones"]["A"] if str(c.value).startswith("=")]
    assert counter == ['=CONCATENATE("Pares decididos: ", COUNTA(Pares!B2:B6), " de 5")']


def test_validate_decisions_with_dedup_levels():
    reg = pd.DataFrame({"N°": [1, 2, 3, 4], "Decisión": [" mismo TRABAJO", "Trabajos distintos", None, "sí"],
                        "Comentario": [None, "x", None, None]})
    dec, problems = validate_decisions(reg, expected_n=[1, 2, 3, 4, 5], levels=DECISIONS)
    assert dec.r2_level.tolist() == [SAME, DIFF]
    assert problems["blank"] == [3] and problems["invalid"] == [[4, "sí"]] and problems["missing_n"] == [5]


@pytest.mark.parametrize("k,n,lo,hi", [
    (5, 10, 0.187086, 0.812914),
    (0, 10, 0.0, 1 - 0.025 ** (1 / 10)),
    (10, 10, 0.025 ** (1 / 10), 1.0),
    (28, 30, 0.779265, 0.991822),
])
def test_clopper_pearson_known_values(k, n, lo, hi):
    ci = clopper_pearson(k, n)
    assert ci == pytest.approx([lo, hi], abs=1e-5)


def test_clopper_pearson_empty():
    assert all(np.isnan(clopper_pearson(0, 0)))


def test_stratum_summary_decisive_and_conservative():
    s = stratum_summary(pd.Series([SAME] * 8 + [DIFF] + [UND]), SAME)
    assert s["decisive"]["k"] == 8 and s["decisive"]["n"] == 9
    assert s["conservative"]["n"] == 10 and s["n_undetermined"] == 1


def test_analyse_and_human_vs_llm():
    df = pd.DataFrame({
        "n": range(1, 13),
        "stratum": ["L1_merge"] * 3 + ["L3_DUPLICATE"] * 4 + ["L3_NOT_DUPLICATE"] * 3 + ["L3_UNCLEAR"] * 2,
        "system_verdict": ["DUPLICATE"] * 7 + ["NOT_DUPLICATE"] * 3 + ["UNCLEAR"] * 2,
        "human": [SAME, SAME, DIFF, SAME, SAME, SAME, UND, DIFF, DIFF, SAME, SAME, DIFF],
    })
    out = analyse(df)
    assert out["strata"]["L1_merge"]["decisive"]["share"] == pytest.approx(2 / 3)
    assert out["strata"]["L3_DUPLICATE"]["decisive"]["n"] == 3
    assert out["strata"]["L3_NOT_DUPLICATE"]["statistic"] == "layer3_npv"
    assert out["strata"]["L3_UNCLEAR"]["human_counts"] == {SAME: 1, DIFF: 1, UND: 0}
    hv = human_vs_llm(df[df.stratum.str.startswith("L3_")])
    # decisive on both sides: DUP x {S,S,S}, NOT x {D,D,S}  -> table [[3,1],[0,2]]
    assert hv["n_compared"] == 6 and hv["n_excluded_human_undetermined"] == 1
    assert hv["n_excluded_llm_unclear"] == 2
    assert hv["raw_agreement"]["share"] == pytest.approx(5 / 6)
    # p_o = 5/6, p_e = (4*3 + 2*3)/36 = 0.5
    assert hv["kappa"] == pytest.approx((5 / 6 - 0.5) / 0.5)
    dis = disagreements(df)
    assert dis.set_index("n").type.to_dict() == {3: "disagreement", 7: "human_undetermined",
                                                 10: "disagreement", 11: "llm_unclear", 12: "llm_unclear"}


def test_raw_records_restore_pre_merge_fields():
    base = dict(source="scielo", source_id="S1", title="T", authors=["A"], year=2020.0, venue="V")
    raw = pd.DataFrame([{**base, "doi": None, "abstract": None},            # repeated key, no DOI
                        {**base, "doi": "10.1/x", "abstract": "abs"},        # repeated key, with DOI
                        {**base, "source_id": "S2", "doi": None, "abstract": None}])
    rr = RawRecords(raw)
    # exact pre-merge copy among repeated keys: kept, not counted as changed
    same = rr.displayed({**base, "doi": "10.1/x", "abstract": "abs", "cluster_id": "c"})
    assert same["doi"] == "10.1/x" and same["cluster_id"] == "c"
    # merged record that gained DOI/abstract from a sibling: pre-merge values restored
    enriched = rr.displayed({**base, "source_id": "S2", "doi": "10.9/y", "abstract": "copied"})
    assert pd.isna(enriched["doi"]) and pd.isna(enriched["abstract"])
    # not found: fallback to the current values
    missing = rr.displayed({**base, "source_id": "S9", "doi": "d"})
    assert missing["doi"] == "d"
    assert rr.fallbacks == ["scielo:S9"] and rr.changed == ["scielo:S2"] and rr.ambiguous == []
    # repeated key, no exact match: first candidate, counted as ambiguous
    rr.displayed({**base, "doi": "10.5/z", "abstract": "other"})
    assert rr.ambiguous == ["scielo:S1"]
