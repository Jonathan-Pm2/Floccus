"""Workbook and agreement helpers of the C03 manual screening (no data files read)."""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
from c03_screening_agreement import PIPE_TO_LEVEL, analyse, disagreements, kappas  # noqa: E402
from make_c03_screening_workbook import INSTRUCTIONS, shuffle_records  # noqa: E402
from make_second_reviewer_workbook import (add_criteria_sheet, add_instructions_sheet,  # noqa: E402
                                           counter_formula, write_records_sheet)
from second_reviewer_agreement import kappa_from_table, validate_decisions  # noqa: E402


def test_shuffle_is_seeded_and_independent_of_input_order():
    ids = pd.DataFrame({"record_id": [f"r{i:03d}" for i in range(50)], "x": range(50)})
    a = shuffle_records(ids, seed=7)
    b = shuffle_records(ids.iloc[::-1], seed=7)
    assert a.record_id.tolist() == b.record_id.tolist()
    assert a.n.tolist() == list(range(1, 51))
    assert a.record_id.tolist() != sorted(a.record_id)
    assert set(a.record_id) == set(ids.record_id)


def test_records_workbook_is_blind_with_dropdowns(tmp_path):
    sample = pd.DataFrame({"n": [1, 2], "record_id": ["r1", "r2"], "title": ["T1", "T2"],
                           "abstract": ["A", None], "keywords": ["k", ""], "venue": ["V", "V"],
                           "year": [2020, None], "doi": ["10.1/x", None],
                           "pipeline_decision": ["INCLUDE", "HUMAN_REVIEW"]})
    wb = Workbook()
    add_instructions_sheet(wb, INSTRUCTIONS, counter_formula("Registros!H2:H3", 2))
    add_criteria_sheet(wb)
    write_records_sheet(wb, sample)
    out = tmp_path / "c03.xlsx"
    wb.save(out)
    wb = load_workbook(out)
    reg = wb["Registros"]
    refs = sorted(str(d.sqref) for d in reg.data_validations.dataValidation)
    assert refs == ["H2:H3", "I2:I3"]
    cells = [c.value for row in reg.iter_rows() for c in row if c.value is not None]
    assert not any(v in ("INCLUDE", "HUMAN_REVIEW", "r1", "r2") for v in cells)
    back = pd.read_excel(out, sheet_name="Registros").astype({"Decisión": object})
    back.loc[0, "Decisión"] = "Excluir"
    dec, problems = validate_decisions(back, expected_n=[1, 2])
    assert dec.n.tolist() == [1] and problems["blank"] == [2]


def frame() -> pd.DataFrame:
    # pipeline INCLUDE x4: human Incluir, Incluir, Dudoso, Excluir
    # pipeline HUMAN_REVIEW x6: human Dudoso x2, Excluir x3, Incluir x1
    return pd.DataFrame({
        "n": range(1, 11),
        "pipeline_decision": ["INCLUDE"] * 4 + ["HUMAN_REVIEW"] * 6,
        "sensitivity_decision": ["INCLUDE"] * 4 + ["HUMAN_REVIEW"] * 4 + ["EXCLUDE"] * 2,
        "human": ["Incluir", "Incluir", "Dudoso", "Excluir", "Dudoso", "Dudoso", "Excluir", "Incluir",
                  "Excluir", "Excluir"],
        "human_reason": [None] * 3 + ["E1 Organismo modelo"] + [None] * 2 + ["E2 Sin modelación formal",
                                                                           None, None, None],
    })


def test_false_positives_and_reasons():
    out = analyse(frame())
    fp = out["false_positives"]["specificity"]
    assert (fp["all_advanced"]["k"], fp["all_advanced"]["n"]) == (4, 10)
    assert (fp["INCLUDE"]["k"], fp["INCLUDE"]["n"]) == (1, 4)
    assert (fp["HUMAN_REVIEW"]["k"], fp["HUMAN_REVIEW"]["n"]) == (3, 6)
    sens = out["false_positives"]["sensitivity"]
    assert (sens["all_advanced"]["k"], sens["all_advanced"]["n"]) == (2, 8)
    assert out["human_counts"] == {"Excluir": 4, "Dudoso": 3, "Incluir": 3}
    assert out["exclusion_reasons"] == {"(sin motivo)": 2, "E1 Organismo modelo": 1,
                                        "E2 Sin modelación formal": 1}


def test_weighted_kappa_with_pipeline_mapping():
    df = frame()
    k = kappas(df.human, df.pipeline_decision.map(PIPE_TO_LEVEL))
    t = k["table_human_rows_pipeline_cols"]
    assert t["Excluir"] == {"Excluir": 0, "Dudoso": 3, "Incluir": 1}
    assert t["Incluir"] == {"Excluir": 0, "Dudoso": 1, "Incluir": 2}
    assert k["raw_agreement"] == pytest.approx(0.4)
    table = np.array([[0, 3, 1], [0, 2, 1], [0, 1, 2]])
    assert k["kappa_linear"] == pytest.approx(kappa_from_table(table, "linear"))
    # unweighted by hand: rows 4,3,3; cols 0,6,4 -> p_e = (0 + 18 + 12) / 100 = 0.30
    assert k["kappa_unweighted"] == pytest.approx((0.4 - 0.30) / 0.70)


def test_disagreements_list_every_mismatch():
    dis = disagreements(frame())
    assert sorted(dis.n) == [3, 4, 7, 8, 9, 10]
