"""Workbook for the second human reviewer of the blind validation sample.

Reviewer 2 (comment 7) asked for a second independent reviewer and a
weighted agreement statistic. This script turns the 465-record stratified
sample into a blind Excel workbook:

- rows are shuffled with a fixed seed, because the sample file is ordered
  by stratum (source x pipeline decision) and that order would reveal the
  pipeline's verdict;
- records get a sequential number; the number -> review_id key is written
  to a separate file that the reviewer must not receive;
- the source column and every pipeline or first-reviewer output are left
  out; the reviewer sees title, abstract, keywords, venue, year and DOI;
- the decision uses an ordinal three-level scale (Include / Unsure /
  Exclude). Collapsing Include + Unsure gives the first reviewer's binary
  scale (both advance to full text), and the three levels can be compared
  with the pipeline's INCLUDE / HUMAN_REVIEW / EXCLUDE by weighted kappa.

Run:  .venv/bin/python scripts/make_second_reviewer_workbook.py --out <file.xlsx>
Out:  <file.xlsx> (to send) and data/processed/validation_reviewer2_key.csv
"""
from __future__ import annotations

import argparse
import re
import uuid
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from biofloc_review.screening.criteria import CRITERIA

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data" / "processed"
SEED = 20261024

DECISIONS = ["Incluir", "Dudoso", "Excluir"]
REASONS = ["E1 Organismo modelo", "E2 Sin modelación formal", "E3 Sistema no acuoso",
           "E4 Coagulación química pura", "No cumple I1 (objeto)", "No cumple I2 (modelo)",
           "No cumple I3 (contexto)", "No cumple I4 (tipo de documento)", "Otro (ver comentario)"]

INSTRUCTIONS = [
    ("Validación ciega del cribado: segundo revisor", "title"),
    ("", None),
    ("Objetivo", "head"),
    ("Decidir, para cada registro, si debe pasar a lectura de texto completo en una revisión "
     "sobre el modelado matemático de flóculos biológicos (biofloc, lodos activados, perifiton, "
     "agregados bacterianos) en sistemas acuáticos.", None),
    ("", None),
    ("Condiciones (importante)", "head"),
    ("1. Trabajo ciego: no consulte las decisiones del algoritmo ni las del otro revisor.", None),
    ("2. Decida solo con lo que ve en la hoja «Registros»: título, resumen, palabras clave, revista, "
     "año y DOI. No busque ni lea el texto completo.", None),
    ("3. Aplique los criterios de la hoja «Criterios», que son los mismos que usaron el algoritmo y "
     "el primer revisor.", None),
    ("4. El orden de los registros es aleatorio y no significa nada.", None),
    ("", None),
    ("Cómo decidir (columna «Decisión», obligatoria)", "head"),
    ("Incluir: cumple I1, I2 e I3, y ninguno de los criterios de exclusión E1–E4.", None),
    ("Excluir: cumple al menos un criterio de exclusión, o claramente no cumple I1, I2 o I3.", None),
    ("Dudoso: la información no alcanza para decidir (resumen muy breve, tipo de modelo poco claro, "
     "contexto mixto). Un registro dudoso pasaría a texto completo.", None),
    ("I4 (tipo de documento): si no se puede verificar con la información disponible, no excluya "
     "por I4.", None),
    ("", None),
    ("Columnas opcionales", "head"),
    ("«Motivo principal de exclusión»: si eligió Excluir, indique el criterio que más pesó.", None),
    ("«Comentario»: cualquier observación útil (por ejemplo, por qué fue dudoso).", None),
    ("", None),
    ("Al terminar", "head"),
    ("Guarde el archivo con el mismo nombre y devuélvalo. Puede trabajar en varias sesiones; la fila "
     "de abajo cuenta cuántos registros lleva decididos.", None),
]


_MOJIBAKE = re.compile("[\u00c2-\u00c3][\u0080-\u00bf]")


def fix_mojibake(text):
    """Undo UTF-8 that was decoded as Latin-1 ('Ã³' -> 'ó'), one character pair at a
    time, so that text mixing broken and correct characters is still repaired."""
    if not isinstance(text, str):
        return text

    def one(m):
        try:
            return m.group(0).encode("latin-1").decode("utf-8")
        except UnicodeDecodeError:
            return m.group(0)
    # a broken non-breaking space whose second byte was already turned into a space
    return _MOJIBAKE.sub(one, text).replace("\u00c2 ", " ")


def _join_keywords(k) -> str:
    if isinstance(k, str):
        return k
    if k is None or len(k) == 0:          # None or an empty array
        return ""
    return "; ".join(str(x) for x in k)


def corpus_text(ids: pd.DataFrame) -> pd.DataFrame:
    """Title, venue, year, DOI, abstract and keywords for the record_id values of
    ids (other columns of ids are kept), with text taken from the corpus and
    mojibake repaired."""
    dec = pd.read_parquet(DATA / "biofloc_final_decisions.parquet",
                          columns=["record_id", "title", "venue", "year", "doi"])
    src = pd.read_parquet(DATA / "biofloc_final_primary_v3.parquet",
                          columns=["source", "source_id", "title_norm", "abstract", "keywords"])
    src["record_id"] = [str(uuid.uuid5(uuid.NAMESPACE_URL, f"{s}:{sid or (tn or '')[:80]}"))
                        for s, sid, tn in zip(src.source, src.source_id, src.title_norm)]
    src["keywords"] = src.keywords.apply(_join_keywords)
    out = (ids.merge(dec, on="record_id", how="left")
              .merge(src[["record_id", "abstract", "keywords"]].drop_duplicates("record_id"),
                     on="record_id", how="left"))
    for col in ("title", "abstract", "keywords", "venue"):
        out[col] = out[col].map(fix_mojibake)
    return out


def load_sample() -> pd.DataFrame:
    """The 465 sampled records, with text taken from the corpus rather than the
    sample CSV, whose text columns were double-encoded on export."""
    ids = pd.read_parquet(DATA / "validation_sample_truth.parquet", columns=["review_id", "record_id"])
    return corpus_text(ids)


# ------------------------------------------------------- workbook helpers
def add_instructions_sheet(wb: Workbook, lines, counter: str):
    """First sheet 'Instrucciones': styled lines plus a bold progress counter."""
    ins = wb.active
    ins.title = "Instrucciones"
    ins.column_dimensions["A"].width = 110
    for i, (text, style) in enumerate(lines, start=1):
        cell = ins.cell(row=i, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if style == "title":
            cell.font = Font(bold=True, size=14)
        elif style == "head":
            cell.font = Font(bold=True, size=11, color="1F4E79")
    ins.cell(row=len(lines) + 2, column=1, value=counter).font = Font(bold=True)
    fit_to_width(ins)
    return ins


def counter_formula(ref: str, total: int, label: str = "Registros decididos") -> str:
    return f'=CONCATENATE("{label}: ", COUNTA({ref}), " de {total}")'


def add_criteria_sheet(wb: Workbook):
    """Sheet 'Criterios' with the screening criteria used by the pipeline."""
    cri = wb.create_sheet("Criterios")
    cri.append(["Código", "Tipo", "Pregunta"])
    kind = {"inclusion": "Inclusión", "exclusion": "Exclusión", "derivation": "Duda (→ Dudoso)"}
    for c in CRITERIA:
        cri.append([c.id, kind.get(str(c.type), str(c.type)), c.question])
    cri.column_dimensions["A"].width = 9
    cri.column_dimensions["B"].width = 16
    cri.column_dimensions["C"].width = 120
    for row in cri.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in cri[1]:
        cell.font = Font(bold=True)
    fit_to_width(cri)
    return cri


def fit_to_width(ws, landscape: bool = True) -> None:
    """Print setup: one page wide, as many pages tall as needed."""
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def add_list_validation(ws, values, ref: str, error: str | None = None) -> DataValidation:
    """Dropdown with the given values on the cell range ref."""
    kw = dict(showErrorMessage=True, errorTitle="Valor no válido", error=error) if error else {}
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True, **kw)
    ws.add_data_validation(dv)
    dv.add(ref)
    return dv


def write_records_sheet(wb: Workbook, sample: pd.DataFrame):
    """Sheet 'Registros' of the screening workbooks: one row per record (columns
    n, title, abstract, keywords, venue, year, doi of sample), decision columns
    H-J with dropdowns."""
    reg = wb.create_sheet("Registros")
    head = ["N°", "Título", "Resumen", "Palabras clave", "Revista", "Año", "DOI",
            "Decisión", "Motivo principal de exclusión", "Comentario"]
    reg.append(head)
    for row in sample.itertuples():
        reg.append([row.n, row.title, row.abstract, row.keywords, row.venue,
                    None if pd.isna(row.year) else int(row.year), row.doi, None, None, None])
    widths = [6, 45, 90, 30, 25, 7, 22, 12, 28, 30]
    for col, w in zip("ABCDEFGHIJ", widths):
        reg.column_dimensions[col].width = w
    fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in reg[1]:
        cell.font = Font(bold=True)
    for row in reg.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in row[7:10]:
            cell.fill = fill
    reg.freeze_panes = "C2"
    reg.print_title_rows = "1:1"
    fit_to_width(reg)
    reg.auto_filter.ref = f"A1:J{len(sample) + 1}"
    last = len(sample) + 1
    add_list_validation(reg, DECISIONS, f"H2:H{last}", "Elija Incluir, Dudoso o Excluir.")
    add_list_validation(reg, REASONS, f"I2:I{last}")
    return reg


def build(out: Path) -> pd.DataFrame:
    sample = load_sample().sort_values("review_id").reset_index(drop=True)
    sample = sample.sample(frac=1, random_state=SEED).reset_index(drop=True)
    sample.insert(0, "n", range(1, len(sample) + 1))
    key = sample[["n", "review_id"]]

    wb = Workbook()
    ins = add_instructions_sheet(wb, INSTRUCTIONS,
                                 counter_formula(f"Registros!H2:H{len(sample) + 1}", len(sample)))
    add_criteria_sheet(wb)
    write_records_sheet(wb, sample)

    wb.active = wb.index(ins)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return key


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, required=True, help="workbook to send to the reviewer")
    args = ap.parse_args()
    key = build(args.out)
    key.to_csv(DATA / "validation_reviewer2_key.csv", index=False)
    print(f"[workbook] {args.out}  records={len(key)}  key -> {DATA / 'validation_reviewer2_key.csv'}")


if __name__ == "__main__":
    main()
